"""Generate synthetic, offline demo material; never read a user's documents.

Usage: python scripts/create_demo.py --output C:/temp/data-search-demo
Requires the project's document dependencies plus reportlab (dev extra).
An existing output directory must be empty to prevent accidental replacement.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
from datetime import datetime
from pathlib import Path


DOCUMENTS = {
    "operations/server_cost_review.md": """# 九月基础设施费用复盘（合成演示资料）

财务发现测试环境在夜间和周末仍然运行，开发服务器平均利用率不足 8%。
团队决定将非生产实例设为工作日 08:30 开机、19:30 关机；周末默认停止。
生产数据库、支付网关和备份节点不参加定时关机，避免影响线上交易。
其次，把过去六个月没有读取的审计附件转移到低频存储，保留检索目录。
预计每月节省 3,200 元。这是内部演示估算，不代表真实项目收益。
负责人是演示运维组，变更编号 INFRA-DEMO-071。上线前先试运行一周。
""",
    "operations/capacity_notes.txt": """容量规划笔记（合成资料）
夜间批量导入任务会把磁盘写满，不应仅凭平均 CPU 低就减少机器数量。
监控显示磁盘 I/O 等待时间升高，查询变慢出现在写入峰值，而不是网络故障。
处理顺序：限制导入并发，错开报表时间，评估 SSD 写入能力。
不要把数据库扫描任务和全量文档索引安排在同一时间段。
""",
    "operations/finops_english.txt": """Cloud commitment review — synthetic example
The production API has a steady baseline load throughout the year.
Before purchasing a reserved capacity commitment, compare one-year utilization,
the cancellation terms, and the portion of spend that remains predictable.
Do not commit seasonal experiments or temporary migration machines.
Renegotiating the vendor contract is separate from shutting down idle test hosts.
Review owner: Demo Procurement. Reference: FINOPS-DEMO-204.
""",
    "orders/refund_policy.md": """# 商城退货说明（合成规则）

商品签收后七个自然日内，可提交退货申请；商品应保持未使用且配件齐全。
质量问题由商家承担回寄运费，个人原因退货的运费由买家承担。
定制刻字商品不适用无理由退货，但质量问题仍可申请售后。
仓库验收通过后，原支付渠道将在三个工作日内发起退款。
退款发起时间与银行实际到账时间不同，客服不要承诺即时到账。
""",
    "orders/return_delivery_case.txt": """售后处理记录（合成资料）
工单 TICKET-DEMO-884：顾客收到杯子时杯柄破裂，包装外箱没有明显破损。
客服已核验照片，批准换货并由商家承担来回运输费用。
顾客不需要先重新购买；仓库收到退件后发送替换商品。
对应演示订单 ORD-DEMO-20260918-0042，问题类型是运输损坏。
""",
    "engineering/resource_scheduler.py": """# Synthetic scheduling example; this file is indexed, never executed.
# 后台索引在内存紧张时让出资源，避免办公软件卡顿。
def should_pause_indexing(available_memory_mb, interactive_query_running):
    if interactive_query_running:
        return True
    return available_memory_mb < 768

# Production databases are excluded from the nightly test-host shutdown plan.
""",
    "engineering/search_settings.yaml": """# 合成配置示例，不会触发服务连接或执行。
node_id: demo-workstation
scan_interval_seconds: 180
resource:
  memory_mb: 1024
  max_disk_mb: 10240
semantic:
  threads: 1
  idle_seconds: 120
database:
  access_mode: read_only
  host: db.example.test
""",
    "engineering/project_plan.json": json.dumps({
        "synthetic": True,
        "project": "资料检索试点",
        "milestones": [
            {"phase": "单机验证", "due": "2026-10-09", "acceptance": "关键词和语义搜索返回可追溯片段"},
            {"phase": "数据库接入", "due": "2026-10-16", "acceptance": "只读查询并限制超时"},
            {"phase": "多机接口", "due": "2026-10-23", "acceptance": "保留节点标识，暂不承诺远程实测"},
        ],
        "risk": "扫描与业务批处理争用磁盘，需要错峰运行",
    }, ensure_ascii=False, indent=2),
    "people/remote_work.html": """<!doctype html><html lang="zh"><head><meta charset="utf-8">
<title>远程办公指引（合成）</title>
<style>.hidden{display:none}</style><script>const ignored = 'DO_NOT_INDEX_SCRIPT';</script>
</head><body><h1>居家工作设备申请</h1>
<p>远程员工可以申请外接显示器和人体工学键盘，审批后由行政寄送。</p>
<p>报修时先登记设备资产编号，并描述现象；无需提供个人账户密码。</p>
<p>新员工第一天先完成设备登记，再参加安全培训。</p></body></html>
""",
    "people/leave_policy_gb18030.txt": """休假申请规则（合成资料）
连续休假超过三个工作日时，应提前五个工作日提交申请，并写好工作交接清单。
突发病假可以先通知直属负责人，恢复后再补齐申请。
这份示例没有定义年假总天数，也不构成真实劳动政策。
""",
}


def _write_texts(root: Path) -> None:
    for relative, text in DOCUMENTS.items():
        path = root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        encoding = "gb18030" if "gb18030" in relative else "utf-8"
        path.write_text(text, encoding=encoding)


def _write_tables(root: Path) -> None:
    orders = root / "orders/orders_export.csv"
    with orders.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerows([
            ["order_id", "customer", "product", "amount_cny", "status", "description"],
            ["ORD-DEMO-20260918-0042", "演示顾客甲", "陶瓷杯", "89.00", "换货处理中", "杯柄运输破损，商家承担回寄费用"],
            ["ORD-DEMO-20260918-0043", "演示顾客乙", "键盘", "329.00", "已签收", "标准配送，无售后申请"],
            ["ORD-DEMO-20260918-0044", "演示顾客丙", "定制刻字本", "69.00", "制作中", "定制商品；尚未发货"],
        ])
    with (root / "engineering/issue_triage.tsv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, delimiter="\t")
        writer.writerows([
            ["issue_id", "component", "priority", "symptom", "action"],
            ["BUG-DEMO-117", "文档解析", "P1", "加密 PDF 无法提取正文", "保留文件名并明确报告需要密码"],
            ["BUG-DEMO-118", "索引更新", "P2", "删除的文件仍出现在搜索结果", "完成目录核对后清除过期条目"],
            ["BUG-DEMO-119", "后台任务", "P2", "电脑内存不足时索引抢占资源", "暂停解析并释放本地向量模型"],
        ])


def _write_office(root: Path) -> None:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    people = root / "people"
    document = Document()
    document.add_heading("新员工入职清单（合成资料）", 0)
    document.add_paragraph("入职前由人事发送材料清单。员工到岗后先核对身份、领取设备，再开通工作账号。")
    document.add_paragraph("第一天完成信息安全培训，不通过聊天发送密码；第二天由导师介绍团队开发流程。")
    document.add_paragraph("试用期第二周安排一次反馈会，由直属负责人确认目标是否清楚。")
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["事项", "时间", "负责组"]):
        cell.text = text
    for values in [("电脑和门禁领取", "第一天上午", "行政"), ("工作账号开通", "设备登记后", "IT 支持"), ("安全培训", "第一天下午", "信息安全组")]:
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    document.save(people / "onboarding_checklist.docx")

    book = Workbook()
    sheet = book.active
    sheet.title = "项目预算"
    for row in [
        ["项目", "月份", "项目类型", "预算金额", "备注"],
        ["资料检索试点", "2026-10", "设备采购", 6000, "只采购一台测试机，复用现有显示器"],
        ["资料检索试点", "2026-10", "培训", 1200, "两次管理员操作培训"],
        ["资料检索试点", "2026-10", "合计", "=SUM(D2:D3)", "公式未计算，不能把缓存缺失当成零"],
    ]:
        sheet.append(row)
    sheet.column_dimensions["E"].width = 52
    schedule = book.create_sheet("里程碑")
    schedule.append(["事项", "计划日期", "说明"])
    schedule.append(["单机验收", datetime(2026, 10, 9), "先证明低资源设备可用"])
    schedule.append(["操作培训", datetime(2026, 10, 12), "管理员学会查看索引未覆盖原因"])
    schedule["B2"].number_format = "yyyy-mm-dd"
    schedule["B3"].number_format = "yyyy-mm-dd"
    book.save(root / "engineering/project_budget.xlsx")

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "资料检索试点评审（合成）"
    slide.placeholders[1].text = "先验证单机安装、正文提取和查询，再保留多机接入接口。\n成功标准包括定位原文、报告未索引原因和遵守资源预算。"
    slide.notes_slide.notes_text_frame.text = "评审备注：不能用文件总 GB 数证明性能。必须同时记录文件数、文本片段数和索引容量。"
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "上线前回退准备"
    slide.placeholders[1].text = "保留上一份可用配置。发现查询结果异常时，先暂停后台索引并保留日志。\n卸载可以选择保留索引，便于重新安装。"
    slide.notes_slide.notes_text_frame.text = "回退演练只使用合成目录，不对正式业务数据库进行写操作。"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(7), Inches(1)).table
    for row, values in enumerate([("检查项", "要求"), ("来源定位", "能回到原文件或数据库记录")]):
        for column, text in enumerate(values):
            table.cell(row, column).text = text
    slides.save(root / "engineering/pilot_review.pptx")


def _write_pdf(root: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = root / "operations/backup_runbook.pdf"
    document = canvas.Canvas(str(path), pagesize=A4)
    document.setTitle("Backup recovery rehearsal — synthetic demo")
    pages = [
        [
            "备份恢复演练手册（合成资料）",
            "每天凌晨两点执行增量备份，每周日执行完整备份。",
            "每月在隔离测试环境中恢复一次备份，检查记录总数和抽样订单。",
            "只看到备份文件存在，不能证明数据可以恢复。",
            "演练编号：DR-DEMO-305。预期恢复时间为四小时。",
        ],
        [
            "Recovery verification checklist",
            "Restore into an isolated test environment, never over the live database.",
            "Compare row counts and a sample of order records after restoration.",
            "Record elapsed recovery time and the most recent recoverable timestamp.",
            "A successful backup job is not proof of a successful recovery.",
        ],
    ]
    for page in pages:
        document.setFont("STSong-Light", 12)
        y = A4[1] - 60
        for line in page:
            document.drawString(40, y, line)
            y -= 30
        document.showPage()
    document.save()


def _write_database(output: Path) -> dict:
    directory = output / "databases"
    directory.mkdir()
    path = directory / "demo.sqlite"
    with sqlite3.connect(path) as connection:
        connection.executescript("""
            CREATE TABLE orders (
                id INTEGER PRIMARY KEY, order_number TEXT UNIQUE NOT NULL,
                customer_name TEXT NOT NULL, amount_cny REAL NOT NULL,
                status TEXT NOT NULL, description TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE employees (
                id INTEGER PRIMARY KEY, employee_code TEXT UNIQUE NOT NULL,
                display_name TEXT NOT NULL, department TEXT NOT NULL,
                onboarding_note TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE support_tickets (
                id INTEGER PRIMARY KEY, ticket_number TEXT UNIQUE NOT NULL,
                order_id INTEGER REFERENCES orders(id), subject TEXT NOT NULL,
                body TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE projects (
                id INTEGER PRIMARY KEY, name TEXT NOT NULL, summary TEXT NOT NULL,
                status TEXT NOT NULL, updated_at TEXT NOT NULL
            );
        """)
        stamp = "2026-09-25T09:00:00Z"
        connection.executemany("INSERT INTO orders VALUES (?, ?, ?, ?, ?, ?, ?)", [
            (1, "ORD-DEMO-20260918-0042", "演示顾客甲", 89.0, "换货处理中", "陶瓷杯运输破损，商家承担回寄费用", stamp),
            (2, "ORD-DEMO-20260918-0043", "演示顾客乙", 329.0, "已签收", "键盘标准配送，无售后申请", stamp),
            (3, "ORD-DEMO-20260918-0044", "演示顾客丙", 69.0, "制作中", "定制刻字本尚未发货", stamp),
        ])
        connection.executemany("INSERT INTO employees VALUES (?, ?, ?, ?, ?, ?)", [
            (1, "EMP-DEMO-001", "演示员工甲", "研发", "已领取电脑，待完成第一天下午的信息安全培训", stamp),
            (2, "EMP-DEMO-002", "演示员工乙", "客户服务", "已完成安全培训，导师安排第二天流程介绍", stamp),
        ])
        connection.executemany("INSERT INTO support_tickets VALUES (?, ?, ?, ?, ?, ?)", [
            (1, "TICKET-DEMO-884", 1, "杯柄运输损坏", "客服已核验照片并批准换货，商家承担运输费用", stamp),
            (2, "TICKET-DEMO-885", None, "后台任务占用资源", "办公电脑可用内存不足时，暂停索引并卸载本地向量模型", stamp),
        ])
        connection.execute("INSERT INTO projects VALUES (?, ?, ?, ?, ?)", (1, "资料检索试点", "先完成单机验证，再扩展多机接入；需要保留来源定位和节点标识", "验证中", stamp))
    config = {
        "id": "demo-sqlite", "kind": "sqlite", "path": str(path.resolve()),
        "allowed_tables": ["orders", "employees", "support_tickets", "projects"],
        "allowed_columns": {
            "orders": ["id", "order_number", "customer_name", "amount_cny", "status", "description", "updated_at"],
            "employees": ["id", "employee_code", "display_name", "department", "onboarding_note", "updated_at"],
            "support_tickets": ["id", "ticket_number", "order_id", "subject", "body", "updated_at"],
            "projects": ["id", "name", "summary", "status", "updated_at"],
        },
        "index": [
            {"table": "orders", "id_column": "id", "text_columns": ["order_number", "description"], "updated_column": "updated_at"},
            {"table": "employees", "id_column": "id", "text_columns": ["employee_code", "onboarding_note"], "updated_column": "updated_at"},
            {"table": "support_tickets", "id_column": "id", "text_columns": ["ticket_number", "subject", "body"], "updated_column": "updated_at"},
            {"table": "projects", "id_column": "id", "text_columns": ["name", "summary"], "updated_column": "updated_at"},
        ],
    }
    (output / "database_config.json").write_text(json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8")
    return config


def create_demo(output: Path) -> dict:
    output = output.expanduser().resolve()
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        raise ValueError("Output must be a new or empty directory; existing files will not be overwritten.")
    # Check optional libraries before writing anything, so a missing dependency
    # does not leave a half-created demonstration dataset.
    import docx  # noqa: F401
    import openpyxl  # noqa: F401
    import pptx  # noqa: F401
    import reportlab  # noqa: F401

    output.mkdir(parents=True, exist_ok=True)
    root = output / "files"
    root.mkdir()
    _write_texts(root)
    _write_tables(root)
    _write_office(root)
    _write_pdf(root)
    _write_database(output)
    files = []
    for path in sorted(root.rglob("*")):
        if path.is_file():
            files.append({
                "source": path.relative_to(output).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            })
    manifest = {
        "schema_version": 1, "synthetic": True,
        "file_root": str(root.resolve()),
        "database_config": str((output / "database_config.json").resolve()),
        "source_count": len(files), "files": files,
        "notes": [
            "All material is synthetic and contains no real business records.",
            "Index only files/ as a file root; attach demo.sqlite through the database adapter.",
            "The workbook intentionally has a formula without a cached result; extraction must not recalculate it.",
            "Reference questions are in tests/evaluation_queries.json in the repository.",
        ],
    }
    (output / "demo_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path, help="New or empty output directory")
    args = parser.parse_args()
    try:
        manifest = create_demo(args.output)
    except (ValueError, ImportError) as exc:
        parser.exit(2, f"{exc}\n")
    print(json.dumps({"output": str(args.output.resolve()), "file_root": manifest["file_root"], "files": manifest["source_count"], "database_config": manifest["database_config"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
