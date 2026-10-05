"""Synthetic, isolated Engine/worker/SQLite benchmark (never scans user files).

Example: python scripts/bench_indexing.py --work .packaging-smoke/indexing-v070
The work directory must not exist. Baseline sources are exported from Git, and
each variant gets its own copied corpus and index. No production daemon starts.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import os
import platform
import shutil
import statistics
import subprocess
import sys
import threading
import time
import zipfile
from collections import Counter
from pathlib import Path

import psutil


class Transactions:
    """Count COMMIT and outer RELEASE separately, without counting triggers."""
    def __init__(self):
        self.counts = Counter()
        self.explicit = False
        self.savepoints = []

    def __call__(self, sql):
        prefix = sql[:24].strip().upper()
        if prefix.startswith("BEGIN"):
            self.explicit = True
            self.counts["begin"] += 1
        elif prefix.startswith("SAVEPOINT "):
            self.savepoints.append(sql.strip().split(None, 1)[1].strip('"`[]'))
            self.counts["savepoint"] += 1
        elif prefix.startswith("RELEASE "):
            name = sql.strip().split()[-1].strip('"`[]')
            if name in self.savepoints:
                index = len(self.savepoints) - 1 - self.savepoints[::-1].index(name)
                if not self.explicit and index == 0:
                    self.counts["outer_savepoint_release"] += 1
                self.savepoints = self.savepoints[:index]
            self.counts["release"] += 1
        elif prefix.startswith("COMMIT"):
            self.counts["commit"] += 1
            self.explicit, self.savepoints = False, []
        elif prefix.startswith("ROLLBACK TO"):
            self.counts["rollback_to"] += 1
            name = sql.strip().split()[-1].strip('"`[]')
            if name in self.savepoints:
                index = len(self.savepoints) - 1 - self.savepoints[::-1].index(name)
                self.savepoints = self.savepoints[:index + 1]
        elif prefix.startswith("ROLLBACK"):
            self.counts["rollback"] += 1
            self.explicit, self.savepoints = False, []

    def report(self):
        return {**dict(self.counts), "durable_transactions": self.counts["commit"] + self.counts["outer_savepoint_release"]}


class ProcessTree:
    """Sample only this isolated Engine and its descendants, not other chats."""
    def __init__(self, interval=.025):
        self.interval, self.stopped = interval, threading.Event()
        self.parent = psutil.Process()
        self.initial, self.latest = {}, {}
        self.peak = self.main_peak = self.children_peak = self.samples = 0
        self.started = time.perf_counter()
        self.sample(initial=True)
        self.thread = threading.Thread(target=self.run, daemon=True)
        self.thread.start()

    def sample(self, initial=False):
        total = children = 0
        try:
            processes = [self.parent] + self.parent.children(recursive=True)
        except psutil.Error:
            processes = [self.parent]
        for process in processes:
            try:
                key = (process.pid, process.create_time())
                cpu = process.cpu_times()
                self.latest[key] = cpu.user + cpu.system
                if initial:
                    self.initial[key] = self.latest[key]
                rss = process.memory_info().rss
                total += rss
                if process.pid == self.parent.pid:
                    self.main_peak = max(self.main_peak, rss)
                else:
                    children += 1
            except psutil.Error:
                pass
        self.peak = max(self.peak, total)
        self.children_peak = max(self.children_peak, children)
        self.samples += 1

    def run(self):
        while not self.stopped.wait(self.interval):
            self.sample()

    def finish(self):
        self.stopped.set()
        self.thread.join()
        self.sample()
        return {"wall_seconds": time.perf_counter() - self.started,
                "cpu_seconds": sum(max(0, value - self.initial.get(key, 0)) for key, value in self.latest.items()),
                "process_tree_rss_peak_bytes": self.peak, "engine_rss_peak_bytes": self.main_peak,
                "descendant_processes_peak": self.children_peak, "resource_samples": self.samples,
                "sample_interval_seconds": self.interval}


def module_from(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def digest(chunks):
    return hashlib.sha256(json.dumps(chunks, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()


def markers(index, version):
    return f"headv{version}f{index:04d}", f"tailf{index:04d}"


def pdf(path, index, version):
    from reportlab.pdfgen.canvas import Canvas
    head, tail = markers(index, version)
    canvas = Canvas(str(path), invariant=1, pageCompression=1)
    for page in range(3):
        for line in range(25):
            text = head if page == line == 0 else tail if page == 2 and line == 24 else "Server indexing budget alpha beta gamma delta 0123456789"
            canvas.drawString(40, 780 - line * 25, text)
        canvas.showPage()
    canvas.save()


def replace_head(path, old, new):
    if path.suffix in {".docx", ".xlsx", ".pptx"}:
        with zipfile.ZipFile(path) as archive:
            members = [(item, archive.read(item)) for item in archive.infolist()]
        with zipfile.ZipFile(path, "w") as archive:
            for item, data in members:
                archive.writestr(item, data.replace(old.encode(), new.encode()))
    else:
        path.write_bytes(path.read_bytes().replace(old.encode(), new.encode()))


def generate(directory, files):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    directory.mkdir(parents=True)
    line = "服务器成本 索引更新 alpha beta gamma delta epsilon budget data 0123456789\n"
    records = []
    for index in range(files):
        kind = ["utf8", "gb18030", "csv", "html", "pdf", "docx", "xlsx", "pptx"][index % 8]
        suffix = {"utf8": "txt", "gb18030": "log"}.get(kind, kind)
        path = directory / f"file-{index:04d}.{suffix}"
        head, tail = markers(index, 0)
        if kind in {"utf8", "gb18030"}:
            prefix = ("ascii prefix data\n" * 4100) if kind == "gb18030" else ""
            path.write_bytes((head + "\n" + prefix + line * 2800 + tail).encode("utf-8" if kind == "utf8" else kind))
        elif kind == "csv":
            with path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow([head, "预算", "备注"])
                for row in range(1700):
                    writer.writerow([f"记录{row}", 1280 + row, line.strip()])
                writer.writerow([tail])
        elif kind == "html":
            path.write_text(f"<html><body><p>{head}</p>" + f"<p>{line}</p>" * 2300 + f"<p>{tail}</p></body></html>", encoding="utf-8")
        elif kind == "pdf":
            pdf(path, index, 0)
        elif kind == "docx":
            document = Document()
            document.add_paragraph(head)
            for _ in range(100):
                document.add_paragraph(line * 3)
            document.add_table(rows=1, cols=1).cell(0, 0).text = tail
            document.save(path)
        elif kind == "xlsx":
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "成本"
            sheet.append([head])
            for row in range(160):
                sheet.append([row, line.strip()])
            sheet.append([tail])
            workbook.save(path)
        else:
            presentation = Presentation()
            for slide_number in range(3):
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = head if slide_number == 0 else "成本预算"
                slide.placeholders[1].text = line * 25
            slide.notes_slide.notes_text_frame.text = tail
            presentation.save(path)
        records.append({"file": path.name, "index": index, "kind": kind, "bytes": path.stat().st_size})
    return records


def mutate(corpus, records):
    for row in records:
        path = corpus / row["file"]
        before = path.stat()
        if row["kind"] == "pdf":
            pdf(path, row["index"], 1)
        else:
            replace_head(path, markers(row["index"], 0)[0], markers(row["index"], 1)[0])
        # Ensure discovery sees an edit even on coarse-time filesystems.
        os.utime(path, ns=(before.st_atime_ns, max(time.time_ns(), before.st_mtime_ns + 2_000_000_000)))


def expected(corpus, records, source):
    extractor = module_from(source / "data_search" / "extractors.py", "reference_extractor")
    chunker = module_from(source / "data_search" / "chunking.py", "reference_chunker")
    result = {}
    for row in records:
        parsed = extractor.extract(str(corpus / row["file"]))
        if parsed["status"] != "ready":
            raise AssertionError((row["file"], parsed["status"], parsed["reason"]))
        chunks = list(chunker.split_chunks(parsed["chunks"]))
        result[row["file"]] = {"chunk_sha256": digest(chunks), "chunks": len(chunks),
                               "extracted_characters": sum(len(chunk["text"]) for chunk in parsed["chunks"])}
    return result


def verify(engine, corpus, records, reference, version):
    documents = engine.store.rows("SELECT id,path,status,reason FROM documents WHERE source_id='files'")
    assert len(documents) == len(records), ("document_count", len(documents), len(records))
    ordered = "ordinal,id" if any(row["name"] == "ordinal" for row in engine.store.rows("PRAGMA table_info(chunks)")) else "id"
    by_name = {Path(row["path"]).name: row for row in documents}
    latencies = []
    for row in records:
        doc = by_name[row["file"]]
        assert doc["status"] == "ready", doc
        chunks = [{"text": chunk["text"], "locator": json.loads(chunk["locator"])} for chunk in
                  engine.store.rows(f"SELECT text,locator FROM chunks WHERE doc_id=? ORDER BY {ordered}", (doc["id"],))]
        assert digest(chunks) == reference[row["file"]]["chunk_sha256"], ("complete_content_or_locator_mismatch", row["file"])
        for token in markers(row["index"], version):
            started = time.perf_counter()
            hits = engine.search(token, "keyword", limit=5)["results"]
            latencies.append(time.perf_counter() - started)
            assert any(Path(hit["path"]).resolve() == (corpus / row["file"]).resolve() for hit in hits), ("missing_keyword", row["file"], token, hits)
        if version:
            assert not engine.search(markers(row["index"], 0)[0], "keyword", limit=5)["results"], ("stale_head", row["file"])
    latencies.sort()
    return {"documents": len(documents), "complete_content_and_locators": True, "keyword_head_and_tail": True,
            "obsolete_keyword_absent": bool(version), "query_count": len(latencies),
            "query_p50_seconds": statistics.median(latencies), "query_p95_seconds": latencies[int((len(latencies) - 1) * .95)]}


def run_phase(engine, name, timeout):
    transactions = Transactions()
    engine.store.db.set_trace_callback(transactions)
    before = engine.telemetry.snapshot({}) if hasattr(engine, "telemetry") else None
    sampler = ProcessTree()
    rounds, errors, full = 0, [], True
    try:
        while True:
            result = engine.scan_once(full=full)
            full = False
            rounds += 1
            if result.get("last_error"):
                errors.append(result["last_error"])
                raise AssertionError((name, errors))
            progress = engine.catalog.progress()
            if progress["retry_files"]:
                raise AssertionError(("unexpected_retry", engine.store.rows("SELECT path,status,reason FROM documents WHERE status NOT IN ('ready','pending')")))
            if not engine._has_work():
                break
            if time.perf_counter() - sampler.started > timeout:
                raise TimeoutError((name, progress))
    finally:
        measured = sampler.finish()
        engine.store.db.set_trace_callback(None)
    performance = None
    if before is not None:
        after = engine.telemetry.snapshot({})
        performance = {
            "stages": {stage: {key: values[key] - before["stages"][stage][key] for key in ("seconds", "calls")}
                       for stage, values in after["stages"].items()},
            "counters": {key: value - before["counters"][key] for key, value in after["counters"].items()},
        }
    return {"phase": name, **measured, "scan_rounds": rounds, "sql_transactions": transactions.report(),
            "catalog": progress, "budget": engine.budget.snapshot(), "parser_limits": engine.parser.control_status,
            "performance": performance}


def worker(args):
    source = args.source.resolve()
    sys.path.insert(0, str(source))
    os.environ["PYTHONPATH"] = str(source)
    from data_search.config import defaults
    from data_search.engine import Engine
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    config = defaults(str(args.data), [str(args.corpus)])
    config["semantic"]["enabled"] = False
    if args.mode == "fixed":
        config["resource"].update(budget_mode="fixed", memory_mb=1024, workers=1,
                                  worker_memory_mb=512, worker_cpu_percent=25, batch_sleep_ms=0)
    report = {"source": str(source), "source_sha256": source_hash(source), "config": config, "phases": []}
    engine = None
    try:
        initialization = ProcessTree()
        engine = Engine(config)
        report["initialization"] = initialization.finish()
        for name, version in [("first", 0), ("unchanged", 0), ("modified", 1)]:
            if version:
                mutate(args.corpus, manifest["files"])
            phase = run_phase(engine, name, args.timeout)
            phase["verification"] = verify(engine, args.corpus, manifest["files"], manifest[f"expected_v{version}"], version)
            report["phases"].append(phase)
            args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            print(json.dumps({"phase": name, "wall_seconds": phase["wall_seconds"], "rss_mib": phase["process_tree_rss_peak_bytes"] / 1048576,
                              "transactions": phase["sql_transactions"]["durable_transactions"]}), flush=True)
    finally:
        if engine:
            engine.close()
    report["completed"] = True
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def source_hash(source):
    hasher = hashlib.sha256()
    for path in sorted(source.rglob("*.py")):
        hasher.update(path.relative_to(source).as_posix().encode())
        hasher.update(path.read_bytes())
    return hasher.hexdigest()


def orchestrate(args):
    repository = Path(__file__).resolve().parents[1]
    work = args.work.resolve()
    work.mkdir(parents=True, exist_ok=False)
    baseline = work / "baseline"
    archive = subprocess.check_output(["git", "archive", "--format=zip", args.baseline, "src"], cwd=repository)
    with zipfile.ZipFile(io.BytesIO(archive)) as bundle:
        for member in bundle.infolist():
            target = (baseline / member.filename).resolve()
            if not target.is_relative_to(baseline.resolve()):
                raise ValueError("unsafe_git_archive_path")
        bundle.extractall(baseline)
    # Freeze current source too: concurrent edits cannot contaminate a run.
    current = work / "current" / "src"
    shutil.copytree(repository / "src", current, ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "*.egg-info"))
    corpus = work / "corpus-v0"
    if args.reuse_corpus:
        previous = args.reuse_corpus.resolve()
        metadata = json.loads((previous / "summary.json").read_text(encoding="utf-8"))
        assert metadata["baseline_source_sha256"] == source_hash(baseline / "src"), "reference source changed"
        manifest = json.loads((previous / "manifest.json").read_text(encoding="utf-8"))
        assert manifest["file_count"] == args.files, "reference file count differs"
        records = manifest["files"]
        shutil.copytree(previous / "corpus-v0", corpus)
    else:
        records = generate(corpus, args.files)
        manifest = {"files": records, "file_count": len(records), "physical_bytes": sum(row["bytes"] for row in records),
                    "by_kind": dict(Counter(row["kind"] for row in records)), "expected_v0": expected(corpus, records, baseline / "src")}
        updated = work / "corpus-v1"
        shutil.copytree(corpus, updated)
        mutate(updated, records)
        manifest["expected_v1"] = expected(updated, records, baseline / "src")
    manifest_path = work / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    summary = {"baseline_commit": subprocess.check_output(["git", "rev-parse", args.baseline], cwd=repository, text=True).strip(),
               "current_base_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=repository, text=True).strip(),
               "baseline_source_sha256": source_hash(baseline / "src"), "current_source_sha256": source_hash(current),
               "platform": platform.platform(), "python": platform.python_version(), "logical_cpus": psutil.cpu_count(),
               "total_memory_bytes": psutil.virtual_memory().total,
               "corpus": {key: manifest[key] for key in ("file_count", "physical_bytes", "by_kind")},
               "semantics": {"semantic_enabled": False, "file_events": "explicit complete reconciliation each phase; no idle scheduler or OS notification delay",
                             "verification": "all indexed chunk text and locators equal baseline extraction+chunking; each file head/tail FTS queries; obsolete head absent after edit",
                             "memory": "25ms sampled Engine plus descendant RSS, including limit helpers; excludes DSH Web/Node/MCP and fixture generation",
                             "cpu": "sampled cumulative parent+child CPU; exited child tail can be undercounted",
                             "sql": "trace-instrumented operation times; durable COMMIT + outer SAVEPOINT RELEASE, schema initialization excluded",
                             "fixed": "both versions 1024MiB process-tree ceiling, 1 parser, 512MiB worker, 25 percent worker CPU, 0 fixed delay",
                             "cache": "no OS disk-cache eviction; initial worker/import cold, later reuse or retirement follows product idle policy; one run per variant"}, "variants": {}}
    if args.reuse_corpus:
        summary["reference_run"] = str(args.reuse_corpus.resolve())
    (work / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"prepared": str(work), "files": len(records), "physical_mib": manifest["physical_bytes"] / 1048576}), flush=True)
    if args.prepare_only:
        return
    for label, source, mode in [("baseline-default", baseline / "src", "default"), ("current-default", current, "default"),
                                ("baseline-fixed", baseline / "src", "fixed"), ("current-fixed", current, "fixed")]:
        if label not in args.variants:
            continue
        directory = work / label
        directory.mkdir()
        variant_corpus = directory / "files"
        shutil.copytree(corpus, variant_corpus)
        output = directory / "report.json"
        environment = os.environ.copy()
        environment["PYTHONPATH"] = str(source)
        with (directory / "run.log").open("w", encoding="utf-8") as log:
            result = subprocess.run([sys.executable, str(Path(__file__).resolve()), "--worker", "--source", str(source),
                                     "--corpus", str(variant_corpus), "--data", str(directory / "index"), "--manifest", str(manifest_path),
                                     "--mode", mode, "--output", str(output), "--timeout", str(args.timeout)],
                                    cwd=directory, env=environment, stdout=log, stderr=subprocess.STDOUT)
        if result.returncode:
            raise RuntimeError(f"{label} failed; inspect {directory / 'run.log'}")
        summary["variants"][label] = json.loads(output.read_text(encoding="utf-8"))
        (work / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"variant": label, "phases": [{"phase": row["phase"], "seconds": row["wall_seconds"]} for row in summary["variants"][label]["phases"]]}), flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--work", type=Path)
    parser.add_argument("--baseline", default="0cb03bb")
    parser.add_argument("--files", type=int, default=384)
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--reuse-corpus", type=Path, help="reuse a prior run's generated corpus and baseline content oracle")
    parser.add_argument("--variants", nargs="+", choices=("baseline-default", "current-default", "baseline-fixed", "current-fixed"),
                        default=["baseline-default", "current-default", "baseline-fixed", "current-fixed"])
    parser.add_argument("--timeout", type=float, default=1200)
    parser.add_argument("--worker", action="store_true")
    for option in ("source", "corpus", "data", "manifest", "output"):
        parser.add_argument("--" + option, type=Path)
    parser.add_argument("--mode", choices=("default", "fixed"), default="default")
    args = parser.parse_args()
    if args.worker:
        worker(args)
    elif not args.work or args.files < 8 or args.files % 8:
        parser.error("--work is required and --files must be a positive multiple of 8")
    else:
        orchestrate(args)


if __name__ == "__main__":
    main()
