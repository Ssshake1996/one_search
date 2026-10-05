"""Small, isolated real-ONNX indexing comparison; never downloads a model.

python scripts/bench_semantic.py --work .packaging-smoke/semantic-v070 \
    --model-dir .runtime/models/bge-small-zh-v1.5

Only fresh synthetic corpus/index directories are written. Both versions use
their real background scheduler, parser children, local BGE model and ANN
publication. The model directory is read-only. This is not a retrieval-quality
evaluation or a large-corpus performance estimate.
"""
from __future__ import annotations

import argparse
from collections import Counter
import io
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
import threading
import time
import zipfile

import psutil

from bench_indexing import ProcessTree, source_hash


def generate(root, count):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen.canvas import Canvas
    root.mkdir(parents=True)
    records = []
    topics = ["服务器内存资源仍有空余时，应允许增加正文解析的并发任务。",
              "数据库备份完成后核对恢复点，避免丢失近期修改的订单记录。",
              "项目会议讨论下周的开发进度，并确认负责人和交付日期。"]
    for index in range(count):
        kind = ['txt', 'log', 'csv', 'html', 'pdf', 'docx', 'xlsx', 'pptx'][index % 8]
        path = root / f'semantic-{index:02d}.{kind}'
        marker = f'firstmarker{index:02d}'
        lines = [f'{marker} 文档{index} 第{row}段。' + topics[index % len(topics)] * 3
                 for row in range(12)]
        if kind in ('txt', 'log'):
            path.write_bytes(('\n'.join(lines)).encode('utf-8' if kind == 'txt' else 'gb18030'))
        elif kind == 'csv':
            path.write_text('marker,description\n' + '\n'.join(f'{marker},{line}' for line in lines), encoding='utf-8')
        elif kind == 'html':
            path.write_text('<html><body>' + ''.join(f'<p>{line}</p>' for line in lines) + '</body></html>', encoding='utf-8')
        elif kind == 'pdf':
            canvas = Canvas(str(path), invariant=1, pageCompression=1)
            for row in range(12):
                canvas.drawString(40, 770 - row * 35, f'{marker} document {index} row {row}: server memory and content indexing budget.')
            canvas.showPage()
            canvas.save()
        elif kind == 'docx':
            document = Document()
            for line in lines:
                document.add_paragraph(line)
            document.save(path)
        elif kind == 'xlsx':
            book = Workbook()
            for row, line in enumerate(lines):
                book.active.append([marker, row, line])
            book.save(path)
        else:
            deck = Presentation()
            for part in range(3):
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = marker
                slide.placeholders[1].text = '\n'.join(lines[part * 4:(part + 1) * 4])
            deck.save(path)
        records.append({'file': path.name, 'kind': kind, 'marker': marker, 'bytes': path.stat().st_size})
    return records


class WorkerCounts:
    def __init__(self):
        self.lock, self.values = threading.Lock(), Counter()

    def install(self):
        from data_search.workers import Worker
        request, start = Worker.request, Worker._start
        owner = self
        def counted_request(worker, message, *args, **kwargs):
            worker._benchmark_method = message['method']
            with owner.lock:
                owner.values[message['method'] + '_requests'] += 1
            return request(worker, message, *args, **kwargs)
        def counted_start(worker):
            with owner.lock:
                owner.values[getattr(worker, '_benchmark_method', 'unknown') + '_process_starts'] += 1
            return start(worker)
        Worker.request, Worker._start = counted_request, counted_start

    def snapshot(self):
        with self.lock:
            return dict(self.values)


def phase(engine, records, counts, trigger, timeout, *, modified=None):
    initial = counts.snapshot()
    sampler = ProcessTree()
    started = time.perf_counter()
    content_seconds = None
    error = None
    try:
        trigger()
        while True:
            rows = engine.store.rows("SELECT name,status,mtime_ns FROM documents WHERE source_id='files'")
            complete = len(rows) == len(records) and all(row['status'] == 'ready' for row in rows)
            if modified:
                complete = complete and any(row['name'] == modified[0] and row['mtime_ns'] == modified[2] for row in rows)
            if content_seconds is None and complete:
                token = modified[1] if modified else records[0]['marker']
                if engine.search(token, 'keyword')['results']:
                    content_seconds = time.perf_counter() - started
            queued = bool(engine.store.rows('SELECT 1 FROM embedding_queue LIMIT 1'))
            vectors = engine.vectors.status()
            if content_seconds is not None and not queued and not vectors['pending'] and not vectors['building']:
                break
            errors = [value for value in (engine.last_error, getattr(engine, 'semantic_error', None), engine.vector_error) if value]
            if errors:
                raise RuntimeError(';'.join(errors))
            if time.perf_counter() - started > timeout:
                raise TimeoutError(json.dumps({'content_seconds': content_seconds, 'documents': rows,
                                              'embedding_queued': queued, 'vectors': vectors}))
            time.sleep(.05)
    except Exception as caught:
        error = {'type': type(caught).__name__, 'message': str(caught)[:1200]}
    finally:
        measured = sampler.finish()
    final = counts.snapshot()
    result = {'content_searchable_seconds': content_seconds, 'all_semantic_published_seconds': measured['wall_seconds'],
              **measured, 'worker_counts': {key: final.get(key, 0) - initial.get(key, 0) for key in final},
              'embedding_rows': engine.store.rows('SELECT count(*) n FROM embeddings')[0]['n'],
              'chunk_rows': engine.store.rows('SELECT count(*) n FROM chunks')[0]['n'],
              'vectors': engine.vectors.status(), 'error': error}
    if error:
        result['completed'] = False
        return result
    # Querying the real encoder + ANN proves availability only, not answer quality.
    query = '电脑有空闲内存，怎样加快文件正文索引' if not modified else '新增段落讨论数据库备份和恢复'
    answer = engine.search(query, 'semantic', limit=5)
    result['semantic_query'] = {'query': query, 'returned_results': len(answer['results']), 'quality_evaluated': False}
    if modified:
        doc = engine.store.rows('SELECT id FROM documents WHERE name=?', (modified[0],))[0]['id']
        chunks = engine.store.rows('SELECT id,hash FROM chunks WHERE doc_id=? AND text LIKE ?', (doc, '%' + modified[1] + '%'))
        assert chunks, 'updated chunk absent'
        from data_search.store import unpack_vector
        vector = unpack_vector(engine.store.rows('SELECT vector FROM embeddings WHERE hash=?', (chunks[0]['hash'],))[0]['vector'])
        # Self-query of the actual new vector verifies published ANN membership;
        # natural-language relevance is deliberately not a pass/fail quality gate.
        hits = engine.vectors.search(vector, 10, filter_sql=' AND d.id=?', filter_args=(doc,))
        result['updated_chunk_ann_self_retrievable'] = chunks[0]['id'] in {hit[0] for hit in hits}
        assert result['updated_chunk_ann_self_retrievable']
    result['completed'] = True
    return result


def worker(args):
    source = args.source.resolve()
    sys.path.insert(0, str(source))
    os.environ['PYTHONPATH'] = str(source)
    from data_search.config import defaults
    from data_search.engine import Engine
    from data_search.model import model_ready, MODEL_ID, SHA256, asset_digest
    config = defaults(str(args.data), [str(args.corpus)])
    config['semantic']['model_dir'] = str(args.model_dir.resolve())
    config['runtime_policy'] = {'enabled': False, 'foreground_grace_seconds': 0}
    report = {'source_sha256': source_hash(source), 'model_id': MODEL_ID, 'config': config, 'phases': []}
    if not model_ready(config['semantic']['model_dir']):
        report.update(completed=False, skipped='pinned_local_model_not_ready')
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        return
    report['model_sha256_verified'] = all(asset_digest(args.model_dir / name) == digest for name, digest in SHA256.items())
    if not report['model_sha256_verified']:
        raise RuntimeError('existing_model_checksum_mismatch_no_download_attempted')
    records = json.loads(args.manifest.read_text(encoding='utf-8'))['files']
    counters = WorkerCounts()
    counters.install()
    engine = Engine(config)
    try:
        first = phase(engine, records, counters, engine.start_background, args.timeout)
        report['phases'].append({'name': 'first', **first})
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        if first['completed']:
            target = args.corpus / records[0]['file']
            text = target.read_text(encoding='utf-8')
            # Change one paragraph, preserving all other paragraphs and files.
            lines = text.splitlines()
            lines[0] = 'updatedsingleneedle 新增段落讨论数据库备份和恢复，其他段落和文件保持原样。'
            target.write_text('\n'.join(lines), encoding='utf-8')
            modified = (target.name, 'updatedsingleneedle', target.stat().st_mtime_ns)
            update = phase(engine, records, counters, engine.scan_event.set, args.timeout, modified=modified)
            report['phases'].append({'name': 'single_paragraph_update', **update})
        report['completed'] = all(row['completed'] for row in report['phases']) and len(report['phases']) == 2
    finally:
        engine.close()
    report['cleanup'] = {'background_threads_stopped': not engine.thread.is_alive() and
                        (not getattr(engine, 'semantic_thread', None) or not engine.semantic_thread.is_alive())}
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def orchestrate(args):
    repository = Path(__file__).resolve().parents[1]
    work = args.work.resolve()
    work.mkdir(parents=True, exist_ok=False)
    baseline = work / 'baseline'
    archive = subprocess.check_output(['git', 'archive', '--format=zip', args.baseline, 'src'], cwd=repository)
    with zipfile.ZipFile(io.BytesIO(archive)) as bundle:
        if any(not (baseline / row.filename).resolve().is_relative_to(baseline.resolve()) for row in bundle.infolist()):
            raise ValueError('unsafe_git_archive_path')
        bundle.extractall(baseline)
    current = work / 'current' / 'src'
    shutil.copytree(repository / 'src', current, ignore=shutil.ignore_patterns('__pycache__', '*.pyc', '*.egg-info'))
    corpus = work / 'corpus'
    records = generate(corpus, args.files)
    manifest = work / 'manifest.json'
    manifest.write_text(json.dumps({'files': records}, ensure_ascii=False, indent=2), encoding='utf-8')
    summary = {'baseline_commit': subprocess.check_output(['git', 'rev-parse', args.baseline], cwd=repository, text=True).strip(),
               'current_source_sha256': source_hash(current), 'model_dir': str(args.model_dir.resolve()),
               'platform': platform.platform(), 'logical_cpus': psutil.cpu_count(), 'total_memory_bytes': psutil.virtual_memory().total,
               'corpus': {'files': len(records), 'bytes': sum(row['bytes'] for row in records), 'by_kind': dict(Counter(row['kind'] for row in records))},
               'method': {'real_background_scheduler': True, 'real_local_onnx': True, 'real_ann': True,
                          'model_downloaded': False, 'file_discovery': 'explicit full reconciliation at phase start; OS watcher may also report the same edit',
                          'resources': 'version defaults; automatic busy/battery and foreground grace pauses disabled in this synthetic test',
                          'rss': '25ms sampled engine plus descendants, excludes DSH host and corpus generation',
                          'timing': 'one run per version, no OS disk cache eviction; initial worker/model cold, update warm',
                          'not_measured': ['retrieval quality', 'large-corpus scaling', 'physical 8GB computer', 'remote server']},
               'variants': {}}
    for name, source in [('baseline', baseline / 'src'), ('current', current)]:
        directory = work / (name + '-run')
        directory.mkdir()
        copied = directory / 'files'
        shutil.copytree(corpus, copied)
        output = directory / 'report.json'
        environment = os.environ.copy()
        environment['PYTHONPATH'] = str(source)
        with (directory / 'run.log').open('w', encoding='utf-8') as log:
            process = subprocess.run([sys.executable, str(Path(__file__).resolve()), '--worker', '--source', str(source),
                '--corpus', str(copied), '--data', str(directory / 'index'), '--manifest', str(manifest),
                '--model-dir', str(args.model_dir.resolve()), '--output', str(output), '--timeout', str(args.timeout)],
                cwd=repository, env=environment, stdout=log, stderr=subprocess.STDOUT, timeout=args.timeout * 2 + 90)
        summary['variants'][name] = {'exit_code': process.returncode,
            'report': json.loads(output.read_text(encoding='utf-8')) if output.is_file() else None}
        (work / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(json.dumps({'variant': name, 'exit_code': process.returncode,
                          'completed': (summary['variants'][name]['report'] or {}).get('completed')}), flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work', type=Path)
    parser.add_argument('--baseline', default='0cb03bb')
    parser.add_argument('--model-dir', type=Path, default=Path('.runtime/models/bge-small-zh-v1.5'))
    parser.add_argument('--files', type=int, default=24, choices=(16, 24, 32))
    parser.add_argument('--timeout', type=float, default=300)
    parser.add_argument('--worker', action='store_true')
    for name in ('source', 'corpus', 'data', 'manifest', 'output'):
        parser.add_argument('--' + name, type=Path)
    args = parser.parse_args()
    if args.worker:
        worker(args)
    else:
        if not args.work:
            parser.error('--work is required')
        orchestrate(args)


if __name__ == '__main__':
    main()
