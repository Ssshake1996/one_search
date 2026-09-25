import codecs
import csv
from datetime import datetime
import zipfile

import pytest

from data_search.extractors import CHUNK_CHARS, extract


def content(result):
    return "".join(chunk["text"] for chunk in result["chunks"])


def assert_contract(result, max_chars=2_000_000):
    assert result["status"] in {"ready", "partial", "error", "encrypted", "unsupported"}
    assert all(0 < len(chunk["text"]) <= CHUNK_CHARS for chunk in result["chunks"])
    assert sum(len(chunk["text"]) for chunk in result["chunks"]) <= max_chars
    assert all(isinstance(chunk["locator"], dict) for chunk in result["chunks"])


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16", "utf-16-be", "gb18030"])
def test_chinese_encodings_and_line_positions(tmp_path, encoding):
    path = tmp_path / "笔记.txt"
    text = "第一行服务器成本\n第二行降低预算\n"
    if encoding == "utf-16-be":
        path.write_bytes(codecs.BOM_UTF16_BE + text.encode(encoding))
    else:
        path.write_bytes(text.encode(encoding))
    result = extract(str(path))
    assert result["status"] == "ready"
    assert content(result) == text
    assert result["chunks"][0]["locator"] == {"line_start": 1, "line_end": 2}
    assert_contract(result)


def test_gb18030_after_large_ascii_prefix(tmp_path):
    path = tmp_path / "late.log"
    path.write_bytes(("a" * 70000 + "中文内容").encode("gb18030"))
    result = extract(str(path))
    assert result["status"] == "ready"
    assert content(result).endswith("中文内容")


@pytest.mark.parametrize("filename", ["a.py", "a.json", "a.yaml", "a.xml", "a.css", "a.sql", "Dockerfile", ".env", "a.tsx"])
def test_code_and_config_are_text_not_executed(tmp_path, filename):
    path = tmp_path / filename
    path.write_text("# 中文配置\nthrow new Error('no execution');", encoding="utf-8")
    result = extract(str(path))
    assert result["status"] == "ready"
    assert "no execution" in content(result)


def test_long_line_is_chunked_and_character_cap_reported(tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("服务器" * 2000 + "\nnext", encoding="utf-8")
    result = extract(str(path), max_chars=2500)
    assert result["status"] == "partial"
    assert result["truncated"]
    assert result["reason"] == "extracted_character_limit"
    assert len(content(result)) == 2500
    assert all(chunk["locator"] == {"line_start": 1, "line_end": 1} for chunk in result["chunks"])
    assert_contract(result, 2500)


def test_crlf_across_block_boundary_and_cr_only(tmp_path):
    path = tmp_path / "line-endings.txt"
    path.write_bytes(b"a" * 8191 + b"\r\nsecond\rthird\r")
    result = extract(str(path))
    assert result["status"] == "ready"
    assert content(result) == "a" * 8191 + "\nsecond\nthird\n"
    assert result["chunks"][-1]["locator"]["line_end"] == 3


def test_truncated_text_locator_ends_at_returned_line(tmp_path):
    path = tmp_path / "lines.txt"
    path.write_text("first\nsecond\nthird\n", encoding="utf-8")
    result = extract(str(path), max_chars=8)
    assert content(result) == "first\nse"
    assert result["chunks"][0]["locator"] == {"line_start": 1, "line_end": 2}


def test_html_ignores_code_and_external_references(tmp_path):
    path = tmp_path / "page.html"
    path.write_text("<html><head><style>SECRET_STYLE</style><script src='https://example.invalid/remote'>SECRET_CODE</script></head>\n<body><h1>中文标题</h1><p>正文 &amp; 成本</p><img src='https://example.invalid/image'><template>SECRET_TEMPLATE</template></body></html>", encoding="utf-8")
    result = extract(str(path))
    assert result["status"] == "ready"
    assert "中文标题正文 & 成本" in content(result)
    assert "SECRET" not in content(result)
    assert result["chunks"][0]["locator"]["line_start"] == 2


@pytest.mark.parametrize("extension,delimiter", [("csv", ","), ("tsv", "\t")])
def test_csv_records_with_quoted_newlines(tmp_path, extension, delimiter):
    path = tmp_path / f"data.{extension}"
    with path.open("w", encoding="gb18030", newline="") as stream:
        writer = csv.writer(stream, delimiter=delimiter)
        writer.writerow(["名称", "备注"])
        writer.writerow(["服务器", "降低成本\n年度预算"])
    result = extract(str(path))
    assert result["status"] == "ready"
    assert "降低成本\n年度预算" in content(result)
    assert result["chunks"][1]["locator"] == {"record": 2, "columns": "1:2"}


def test_csv_giant_record_stops_with_reason(tmp_path):
    path = tmp_path / "large.csv"
    path.write_text("header\n\"" + "x\n" * 700 + "\"", encoding="utf-8")
    result = extract(str(path), max_chars=100)
    assert result["status"] == "partial"
    assert result["reason"] == "csv_record_limit"
    assert content(result) == "header"


def test_docx_body_tables_and_no_fake_page_locator(tmp_path):
    from docx import Document

    path = tmp_path / "document.docx"
    document = Document()
    document.add_paragraph("降低服务器成本")
    document.add_paragraph("第二段资料")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表格预算"
    document.save(path)
    result = extract(str(path))
    assert result["status"] == "ready"
    assert "降低服务器成本" in content(result)
    assert "表格预算" in content(result)
    assert result["chunks"][0]["locator"]["paragraph"] == 1
    assert result["chunks"][-1]["locator"]["table"] == 1
    assert not any("page" in chunk["locator"] for chunk in result["chunks"])
    assert_contract(result)


def test_xlsx_formula_cache_cells_and_sheet_names(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "data.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "成本预算"
    sheet["A1"] = "服务器"
    sheet["B2"] = "=1+2"
    sheet["C3"] = "=2+3"
    sheet["D4"] = datetime(2026, 9, 25)
    sheet["E5"] = "=A1+1"
    sheet["E6"] = "=A2+1"
    book.save(path)
    with zipfile.ZipFile(path) as archive:
        members = {info.filename: archive.read(info) for info in archive.infolist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(b"<f>1+2</f><v></v>", b"<f>1+2</f><v>3</v>")
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(b"<f>A1+1</f>", b'<f t="shared" si="0" ref="E5:E6">A1+1</f>').replace(b"<f>A2+1</f>", b'<f t="shared" si="0"/>')
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    result = extract(str(path))
    assert result["status"] == "ready"
    assert "成本预算" in content(result)
    cells = {chunk["locator"]["cells"]: chunk for chunk in result["chunks"]}
    assert cells["A1"]["text"] == "服务器"
    assert cells["B2"]["text"] == "=1+2\n[cached] 3"
    assert cells["B2"]["locator"]["cached_value_present"] is True
    assert cells["C3"]["locator"]["cached_value_present"] is False
    assert "unavailable" in cells["C3"]["text"]
    assert cells["D4"]["text"] == "2026-09-25T00:00:00"
    assert cells["E6"]["text"].startswith("=A2+1\n")
    assert_contract(result)


def test_pptx_slide_text_tables_and_notes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "成本优化"
    slide.placeholders[1].text = "降低服务器闲置资源"
    slide.notes_slide.notes_text_frame.text = "演讲备注内容"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "表格成本"
    presentation.save(path)
    result = extract(str(path))
    assert result["status"] == "ready"
    assert "成本优化" in content(result)
    assert "表格成本" in content(result)
    notes = [chunk for chunk in result["chunks"] if chunk["locator"]["notes"]]
    assert any("演讲备注内容" in chunk["text"] for chunk in notes)
    assert all(chunk["locator"]["slide"] == 1 for chunk in result["chunks"])


def make_pdf(path):
    from reportlab.pdfgen.canvas import Canvas

    canvas = Canvas(str(path))
    canvas.drawString(50, 700, "First page server costs")
    canvas.showPage()
    canvas.drawString(50, 700, "Second page budget")
    canvas.save()


def test_pdf_text_pages_and_character_limit(tmp_path):
    path = tmp_path / "text.pdf"
    make_pdf(path)
    result = extract(str(path))
    assert result["status"] == "ready"
    assert "First page" in content(result)
    assert "Second page" in content(result)
    assert {chunk["locator"]["page"] for chunk in result["chunks"]} == {1, 2}
    partial = extract(str(path), max_chars=8)
    assert partial["status"] == "partial"
    assert content(partial) == "First pa"
    assert_contract(partial, 8)


def test_encrypted_pdf_and_office_reported(tmp_path):
    from pypdf import PdfReader, PdfWriter

    source = tmp_path / "source.pdf"
    make_pdf(source)
    reader = PdfReader(source)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("secret")
    encrypted = tmp_path / "encrypted.pdf"
    with encrypted.open("wb") as stream:
        writer.write(stream)
    assert extract(str(encrypted))["status"] == "encrypted"
    # Minimal compound-container signature used only to test identification;
    # decrypting actual Office documents is outside the supported scope.
    office = tmp_path / "encrypted.docx"
    office.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + "EncryptedPackage".encode("utf-16-le"))
    assert extract(str(office))["status"] == "encrypted"


def test_pdf_without_text_reports_ocr_unavailable(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    with path.open("wb") as stream:
        writer.write(stream)
    result = extract(str(path))
    assert result["status"] == "unsupported"
    assert "OCR" in result["reason"]


@pytest.mark.parametrize("suffix", ["docx", "xlsx", "pptx", "pdf"])
def test_damaged_documents_report_error(tmp_path, suffix):
    path = tmp_path / f"broken.{suffix}"
    path.write_bytes(b"this is not a document")
    result = extract(str(path))
    assert result["status"] == "error"
    assert result["reason"]
    assert not result["chunks"]


def test_archive_expansion_and_xml_entity_budgets(tmp_path):
    large = tmp_path / "large.docx"
    with zipfile.ZipFile(large, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * (8 * 1024 * 1024 + 1))
    result = extract(str(large), max_chars=100)
    assert result["status"] == "partial"
    assert result["reason"] == "archive_expansion_limit"
    entity = tmp_path / "entity.docx"
    with zipfile.ZipFile(entity, "w") as archive:
        archive.writestr("word/document.xml", '<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "hello">]><a>&x;</a>')
    result = extract(str(entity))
    assert result["status"] == "error"
    assert "DTDForbidden" in result["reason"]


def test_office_character_limit_reports_partial(tmp_path):
    from docx import Document

    path = tmp_path / "long.docx"
    document = Document()
    document.add_paragraph("中文内容" * 1000)
    document.save(path)
    result = extract(str(path), max_chars=500)
    assert result["status"] == "partial"
    assert len(content(result)) == 500
    assert_contract(result, 500)


def test_non_supported_format_and_bad_budget(tmp_path):
    path = tmp_path / "image.png"
    path.write_bytes(b"image")
    assert extract(str(path))["status"] == "unsupported"
    assert extract(str(path), max_chars=0)["status"] == "error"


def test_binary_disguised_as_text_is_rejected(tmp_path):
    path = tmp_path / "binary.txt"
    path.write_bytes(b"\0\x01\0\x02\0\x03\0\x04\0\x05\0\x06\0\x07")
    result = extract(str(path))
    # It is structurally possible UTF-16, so a BOM is needed before accepting
    # this ambiguous all-control-character sample as a text file.
    assert result["status"] == "error"
