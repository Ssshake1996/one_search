"""Reproducible extractor-only benchmark; generates no user-file fixtures.

Run with ``python tests/benchmark_extractors.py --baseline old_extractors.py
--output report.json``. Wall time excludes fixture creation and imports. The
separate allocation measurement is Python tracemalloc, not process RSS. This
does not measure worker startup, indexing, semantic work, or directory scans.
"""

from __future__ import annotations

import argparse
import gc
import hashlib
import importlib.util
import json
import platform
import statistics
import tempfile
import time
import tracemalloc
import zipfile
from pathlib import Path

from data_search import extractors


def fixtures(directory: Path) -> list[Path]:
    text = "服务器成本和资料检索 mixed text 0123456789\r\n" * 40_000
    (directory / "utf8.txt").write_bytes(text.encode("utf-8"))
    (directory / "legacy.log").write_bytes(("a" * 70_000 + text).encode("gb18030"))
    (directory / "rows.csv").write_text("名称,预算,备注\n" + "服务器,1280,年度费用\n" * 8000, encoding="utf-8")
    (directory / "page.html").write_text("<html><body>" + "<p>服务器预算 mixed text</p>" * 6000 + "</body></html>", encoding="utf-8")
    word_ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    with zipfile.ZipFile(directory / "paragraphs.docx", "w", zipfile.ZIP_DEFLATED) as archive:
        paragraphs = ''.join(f'<w:p><w:pPr><w:spacing w:after="120"/></w:pPr><w:r><w:rPr><w:sz w:val="22"/></w:rPr><w:t>第{i}段 服务器预算 mixed text</w:t></w:r></w:p>' for i in range(8000))
        archive.writestr("word/document.xml", f'<w:document {word_ns}><w:body>{paragraphs}</w:body></w:document>')
    with zipfile.ZipFile(directory / "large-paragraphs.docx", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", f'<w:document {word_ns}><w:body>{paragraphs}{paragraphs}</w:body></w:document>')
    sheet_ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    with zipfile.ZipFile(directory / "cells.xlsx", "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", f'<workbook {sheet_ns} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="成本" r:id="one"/></sheets></workbook>')
        archive.writestr("xl/_rels/workbook.xml.rels", '<Relationships><Relationship Id="one" Target="worksheets/sheet1.xml"/></Relationships>')
        rows = ''.join(f'<row r="{i}"><c r="A{i}" t="inlineStr"><is><t>第{i}行 服务器预算 mixed text</t></is></c></row>' for i in range(1, 12_001))
        archive.writestr("xl/worksheets/sheet1.xml", f'<worksheet {sheet_ns}><sheetData>{rows}</sheetData></worksheet>')
    with zipfile.ZipFile(directory / "cells.xlsx") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(directory / "large-cells.xlsx", "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            if name.endswith("sheet1.xml"):
                data = data.replace(b"<is><t>", b"<is><r><t>").replace(b"</t></is>", b"</t></r>" + b"<r><rPr><b/><sz val=\"11\"/></rPr><t>budget rich text </t></r>" * 6 + b"</is>")
            archive.writestr(name, data)
    from pptx import Presentation
    presentation = Presentation()
    for i in range(15):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"预算 {i}"
        slide.placeholders[1].text = "服务器预算 mixed text\n" * 20
        slide.notes_slide.notes_text_frame.text = "备注"
    presentation.save(directory / "slides.pptx")
    from reportlab.pdfgen.canvas import Canvas
    canvas = Canvas(str(directory / "pages.pdf"))
    for _ in range(15):
        for line in range(30):
            canvas.drawString(50, 780 - line * 20, "Server budget and indexing mixed text")
        canvas.showPage()
    canvas.save()
    return sorted(directory.iterdir())


def signature(result):
    return hashlib.sha256(json.dumps(result, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def measure(module, path, repetitions):
    warm = module.extract(str(path))
    expected = signature(warm)
    times = []
    for _ in range(repetitions):
        gc.collect()
        start = time.perf_counter()
        result = module.extract(str(path))
        times.append(time.perf_counter() - start)
        assert signature(result) == expected
        del result
    gc.collect()
    tracemalloc.start()
    result = module.extract(str(path))
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert signature(result) == expected
    return {
        "median_seconds": statistics.median(times), "runs_seconds": times,
        "python_allocation_peak_bytes": peak, "result_sha256": expected,
        "status": result["status"], "chunks": len(result["chunks"]),
        "characters": sum(len(chunk["text"]) for chunk in result["chunks"]),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--repetitions", type=int, default=5)
    args = parser.parse_args()
    if args.repetitions < 1:
        parser.error("repetitions must be positive")
    spec = importlib.util.spec_from_file_location("baseline_extractors", args.baseline)
    baseline = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(baseline)
    report = {
        "platform": platform.platform(), "python": platform.python_version(),
        "scope": "synthetic extractor-only warm runs; no user files, subprocess startup, index writes, or semantic inference",
        "memory_metric": "separate tracemalloc run; Python allocation peak, not RSS",
        "baseline_sha256": hashlib.sha256(args.baseline.read_bytes()).hexdigest(),
        "current_sha256": hashlib.sha256(Path(extractors.__file__).read_bytes()).hexdigest(),
        "fixtures": [],
    }
    with tempfile.TemporaryDirectory(prefix="one-search-extractors-") as temporary:
        for path in fixtures(Path(temporary)):
            before = measure(baseline, path, args.repetitions)
            after = measure(extractors, path, args.repetitions)
            assert before["result_sha256"] == after["result_sha256"], path.name
            report["fixtures"].append({
                "file": path.name, "bytes": path.stat().st_size,
                "baseline": before, "current": after,
                "speedup": before["median_seconds"] / after["median_seconds"],
                "peak_allocation_ratio": after["python_allocation_peak_bytes"] / before["python_allocation_peak_bytes"],
            })
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps([{ "file": row["file"], "speedup": round(row["speedup"], 2), "peak_ratio": round(row["peak_allocation_ratio"], 2)} for row in report["fixtures"]], ensure_ascii=False))


if __name__ == "__main__":
    main()
